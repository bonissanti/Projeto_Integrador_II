from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import BinaryIO, TypedDict
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from django.db import transaction

from .models import HairStock


class StockWorkbookError(ValueError):
    pass


class StockItem(TypedDict):
    name: str
    unit: str
    minimum_stock: Decimal
    initial_quantity: Decimal
    unit_price: Decimal | None


class ConsumptionRow(TypedDict):
    procedure: str
    product: str
    unit: str
    minimum: Decimal
    average: Decimal
    maximum: Decimal


@dataclass(frozen=True)
class StockWorkbookData:
    stock_items: list[StockItem]
    consumption_rows: list[ConsumptionRow]


@dataclass(frozen=True)
class StockImportResult:
    created_count: int
    skipped_names: list[str]


def _read_decimal(value: object, field_name: str, row_number: int) -> Decimal:
    if value is None or value == "":
        raise StockWorkbookError(
            f"Linha {row_number}: o campo “{field_name}” está vazio."
        )

    try:
        number = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, AttributeError):
        raise StockWorkbookError(
            f"Linha {row_number}: o campo “{field_name}” precisa ser numérico."
        ) from None

    if not number.is_finite() or number < 0:
        raise StockWorkbookError(
            f"Linha {row_number}: o campo “{field_name}” precisa ser zero ou positivo."
        )

    return number


def _header_columns(sheet, required_headers: tuple[str, ...]) -> dict[str, int]:
    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(sheet[1])
        if cell.value is not None
    }
    missing = [header for header in required_headers if header not in headers]
    if missing:
        raise StockWorkbookError(
            f"A aba “{sheet.title}” não contém as colunas obrigatórias: "
            f"{', '.join(missing)}."
        )
    return headers


def _parse_stock_items(sheet) -> list[StockItem]:
    required = (
        "Produto",
        "Unidade de compra",
        "Estoque mínimo",
        "Estoque inicial",
    )
    columns = _header_columns(sheet, required)
    seen_names = set()
    items: list[StockItem] = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = {header: row[index] if index < len(row) else None for header, index in columns.items()}
        if all(value is None or value == "" for value in values.values()):
            continue

        name = str(values["Produto"] or "").strip()
        unit = str(values["Unidade de compra"] or "").strip()
        if not name or not unit:
            raise StockWorkbookError(
                f"Linha {row_number}: produto e unidade de compra são obrigatórios."
            )
        if len(name) > 150 or len(unit) > 20:
            raise StockWorkbookError(
                f"Linha {row_number}: produto ou unidade excede o limite aceito pelo sistema."
            )

        normalized_name = name.casefold()
        if normalized_name in seen_names:
            raise StockWorkbookError(
                f"Linha {row_number}: o produto “{name}” aparece mais de uma vez."
            )
        seen_names.add(normalized_name)

        price_column = columns.get("Preço unitário")
        raw_price = row[price_column] if price_column is not None and price_column < len(row) else None
        price = (
            _read_decimal(raw_price, "Preço unitário", row_number)
            if raw_price is not None and raw_price != ""
            else None
        )
        minimum_stock = _read_decimal(
            values["Estoque mínimo"], "Estoque mínimo", row_number
        )
        initial_quantity = _read_decimal(
            values["Estoque inicial"], "Estoque inicial", row_number
        )
        if minimum_stock > Decimal("9999.999999") or initial_quantity > Decimal("9999.999999"):
            raise StockWorkbookError(
                f"Linha {row_number}: o estoque não pode exceder 9.999,999999 unidades."
            )
        if price is not None and price > Decimal("999999.99"):
            raise StockWorkbookError(
                f"Linha {row_number}: o preço unitário excede o limite aceito pelo sistema."
            )

        items.append({
            "name": name,
            "unit": unit,
            "minimum_stock": minimum_stock,
            "initial_quantity": initial_quantity,
            "unit_price": price,
        })

    if not items:
        raise StockWorkbookError("A aba “Estoque inicial” não tem itens para importar.")
    return items


def _parse_consumption_rows(sheet) -> list[ConsumptionRow]:
    required = (
        "Procedimento",
        "Produto",
        "Unidade de consumo",
        "Consumo mínimo",
        "Consumo médio",
        "Consumo máximo",
    )
    columns = _header_columns(sheet, required)
    rows: list[ConsumptionRow] = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = {header: row[index] if index < len(row) else None for header, index in columns.items()}
        if all(value is None or value == "" for value in values.values()):
            continue

        procedure = str(values["Procedimento"] or "").strip()
        product = str(values["Produto"] or "").strip()
        unit = str(values["Unidade de consumo"] or "").strip()
        if not procedure or not product or not unit:
            raise StockWorkbookError(
                f"Linha {row_number} da aba “{sheet.title}”: procedimento, produto e unidade são obrigatórios."
            )
        if len(unit) > 20:
            raise StockWorkbookError(
                f"Linha {row_number} da aba “{sheet.title}”: a unidade excede 20 caracteres."
            )

        minimum = _read_decimal(
            values["Consumo mínimo"], "Consumo mínimo", row_number
        )
        average = _read_decimal(
            values["Consumo médio"], "Consumo médio", row_number
        )
        maximum = _read_decimal(
            values["Consumo máximo"], "Consumo máximo", row_number
        )
        if not minimum <= average <= maximum:
            raise StockWorkbookError(
                f"Linha {row_number} da aba “{sheet.title}”: os consumos devem "
                "seguir a ordem mínimo, médio e máximo."
            )

        rows.append({
            "procedure": procedure,
            "product": product,
            "unit": unit,
            "minimum": minimum,
            "average": average,
            "maximum": maximum,
        })

    if not rows:
        raise StockWorkbookError("A aba “Consumo por procedimento” não tem dados.")
    return rows


def parse_stock_workbook(file_obj: BinaryIO) -> StockWorkbookData:
    try:
        workbook = load_workbook(file_obj, data_only=True, read_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise StockWorkbookError(
            "Não foi possível ler o arquivo. Envie uma planilha Excel .xlsx válida."
        ) from exc

    try:
        required_sheets = {"Estoque inicial", "Consumo por procedimento"}
        missing_sheets = required_sheets.difference(workbook.sheetnames)
        if missing_sheets:
            raise StockWorkbookError(
                "A planilha precisa conter as abas “Estoque inicial” e "
                "“Consumo por procedimento”."
            )

        return StockWorkbookData(
            stock_items=_parse_stock_items(workbook["Estoque inicial"]),
            consumption_rows=_parse_consumption_rows(
                workbook["Consumo por procedimento"]
            ),
        )
    finally:
        workbook.close()


@transaction.atomic
def import_stock_items(
    stock_items: list[StockItem],
    confirm_zero_prices: bool,
) -> StockImportResult:
    existing_names = {
        name.casefold()
        for name in HairStock.active_objects.filter(deleted=False).values_list(
            "name", flat=True
        )
    }
    new_items = [
        item for item in stock_items
        if item["name"].casefold() not in existing_names
    ]
    missing_prices = [item for item in new_items if item["unit_price"] is None]
    if missing_prices and not confirm_zero_prices:
        raise StockWorkbookError(
            "Confirme que os itens sem preço serão importados com preço unitário 0,00."
        )

    created_count = 0
    skipped_names = []

    for item in stock_items:
        if HairStock.active_objects.filter(
            name__iexact=item["name"],
            deleted=False,
        ).exists():
            skipped_names.append(item["name"])
            continue

        unit_price = item["unit_price"] or Decimal("0.00")
        hair_stock = HairStock.objects.create(
            name=item["name"],
            unit=item["unit"],
            unit_price=unit_price,
            quantity_on_hand=Decimal("0"),
            minimum_stock=item["minimum_stock"],
        )
        if item["initial_quantity"] > 0:
            HairStock.objects.registrar_entrada(
                hair_stock.id,
                item["initial_quantity"],
                unit_price=unit_price,
                note="Estoque inicial importado da planilha.",
            )
        created_count += 1

    return StockImportResult(
        created_count=created_count,
        skipped_names=skipped_names,
    )
