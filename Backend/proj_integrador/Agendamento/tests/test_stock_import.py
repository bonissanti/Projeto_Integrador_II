from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Permission
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import Workbook

from Agendamento.models import HairMovement, HairStock
from Agendamento.stock_import import (
    StockWorkbookError,
    import_stock_items,
    parse_stock_workbook,
)


class TestLeituraDaPlanilhaDeEstoque(SimpleTestCase):
    def criar_planilha(self, inclui_preco=False):
        workbook = Workbook()
        estoque = workbook.active
        estoque.title = "Estoque inicial"
        headers = [
            "Categoria",
            "Produto",
            "Unidade de compra",
            "Estoque mínimo",
            "Estoque inicial",
            "Observação",
        ]
        if inclui_preco:
            headers.append("Preço unitário")
        estoque.append(headers)
        row = ["Finalização", "Mousse fixadora 200 ml", "frasco", 3, 6, "Alta rotatividade"]
        if inclui_preco:
            row.append(12.50)
        estoque.append(row)

        consumo = workbook.create_sheet("Consumo por procedimento")
        consumo.append([
            "Procedimento",
            "Categoria",
            "Produto",
            "Unidade de consumo",
            "Consumo mínimo",
            "Consumo médio",
            "Consumo máximo",
            "Observações",
        ])
        consumo.append(["Box braids", "Finalização", "Mousse fixadora", "ml", 20, 30, 40, ""])

        file = BytesIO()
        workbook.save(file)
        file.seek(0)
        return file

    def test_le_planilha_original_e_preserva_preco_ausente(self):
        planilha = parse_stock_workbook(self.criar_planilha())

        self.assertEqual(len(planilha.stock_items), 1)
        self.assertEqual(planilha.stock_items[0]["name"], "Mousse fixadora 200 ml")
        self.assertEqual(planilha.stock_items[0]["unit"], "frasco")
        self.assertEqual(planilha.stock_items[0]["minimum_stock"], Decimal("3"))
        self.assertEqual(planilha.stock_items[0]["initial_quantity"], Decimal("6"))
        self.assertIsNone(planilha.stock_items[0]["unit_price"])
        self.assertEqual(len(planilha.consumption_rows), 1)
        self.assertEqual(planilha.consumption_rows[0]["average"], Decimal("30"))

    def test_le_preco_quando_a_planilha_tiver_a_coluna(self):
        planilha = parse_stock_workbook(self.criar_planilha(inclui_preco=True))

        self.assertEqual(planilha.stock_items[0]["unit_price"], Decimal("12.5"))

    def test_rejeita_planilha_sem_aba_de_estoque(self):
        workbook = Workbook()
        workbook.active.title = "Outra aba"
        file = BytesIO()
        workbook.save(file)
        file.seek(0)

        with self.assertRaises(StockWorkbookError):
            parse_stock_workbook(file)


class TestImportacaoDoEstoque(TestCase):
    def setUp(self):
        self.item = {
            "name": "Mousse fixadora 200 ml",
            "unit": "frasco",
            "minimum_stock": Decimal("3"),
            "initial_quantity": Decimal("6"),
            "unit_price": None,
        }

    def test_exige_confirmacao_para_preco_zero_e_preserva_historico(self):
        with self.assertRaises(StockWorkbookError):
            import_stock_items([self.item], confirm_zero_prices=False)

        self.assertFalse(HairStock.objects.exists())

        resultado = import_stock_items([self.item], confirm_zero_prices=True)

        estoque = HairStock.objects.get(name=self.item["name"])
        movimento = HairMovement.objects.get(hair_stock=estoque)
        self.assertEqual(resultado.created_count, 1)
        self.assertEqual(resultado.skipped_names, [])
        self.assertEqual(estoque.unit_price, Decimal("0.00"))
        self.assertEqual(estoque.quantity_on_hand, Decimal("6.000000"))
        self.assertEqual(estoque.minimum_stock, Decimal("3.000000"))
        self.assertEqual(movimento.quantity, Decimal("6.000000"))
        self.assertEqual(movimento.unit_price_at_time, Decimal("0.00"))
        self.assertEqual(movimento.note, "Estoque inicial importado da planilha.")

    def test_reimportacao_nao_duplica_item_nem_soma_saldo(self):
        primeiro = import_stock_items([self.item], confirm_zero_prices=True)
        segundo = import_stock_items([self.item], confirm_zero_prices=False)

        self.assertEqual(primeiro.created_count, 1)
        self.assertEqual(segundo.created_count, 0)
        self.assertEqual(segundo.skipped_names, [self.item["name"]])
        self.assertEqual(HairStock.objects.count(), 1)
        self.assertEqual(HairMovement.objects.count(), 1)

    def test_importacao_no_admin_exige_confirmar_preco_zero(self):
        usuario = get_user_model().objects.create_superuser(
            username="admin-importacao",
            email="importacao@example.com",
            password="senha-teste",
        )
        self.client.force_login(usuario)
        arquivo = self.criar_planilha_upload()
        url = reverse("admin:Agendamento_hairstock_importar")

        resposta = self.client.post(url, {"arquivo": arquivo, "action": "upload"})

        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "não informa o preço unitário")
        self.assertContains(resposta, "0,00")
        self.assertContains(resposta, "stock_import.css")
        self.assertContains(resposta, "Consumo por procedimento")
        self.assertContains(resposta, "confirm")
        self.assertFalse(HairStock.objects.exists())

        resposta = self.client.post(url, {"action": "confirm"})
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Confirme que os itens sem preço")
        self.assertFalse(HairStock.objects.exists())

        resposta = self.client.post(
            url,
            {"action": "confirm", "confirmar_preco_zero": "on"},
        )

        self.assertRedirects(
            resposta,
            reverse("admin:Agendamento_hairstock_changelist"),
        )
        self.assertEqual(HairStock.objects.get().unit_price, Decimal("0.00"))
        self.assertEqual(HairMovement.objects.count(), 1)

    def test_importacao_no_admin_exige_permissao_de_adicao(self):
        usuario = get_user_model().objects.create_user(
            username="leitor-estoque",
            password="senha-teste",
            is_staff=True,
        )
        permissao = Permission.objects.get(
            content_type__app_label="Agendamento",
            codename="view_hairstock",
        )
        usuario.user_permissions.add(permissao)
        self.client.force_login(usuario)

        resposta = self.client.get(reverse("admin:Agendamento_hairstock_importar"))

        self.assertEqual(resposta.status_code, 403)

    @staticmethod
    def criar_planilha_upload():
        workbook = Workbook()
        estoque = workbook.active
        estoque.title = "Estoque inicial"
        estoque.append([
            "Categoria",
            "Produto",
            "Unidade de compra",
            "Estoque mínimo",
            "Estoque inicial",
            "Observação",
        ])
        estoque.append(["Finalização", "Mousse fixadora 200 ml", "frasco", 3, 6, ""])
        consumo = workbook.create_sheet("Consumo por procedimento")
        consumo.append([
            "Procedimento",
            "Categoria",
            "Produto",
            "Unidade de consumo",
            "Consumo mínimo",
            "Consumo médio",
            "Consumo máximo",
            "Observações",
        ])
        consumo.append(["Box braids", "Finalização", "Mousse fixadora", "ml", 20, 30, 40, ""])
        file = BytesIO()
        workbook.save(file)
        return SimpleUploadedFile(
            "estoque.xlsx",
            file.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_rejeita_saldo_negativo(self):
        workbook = Workbook()
        estoque = workbook.active
        estoque.title = "Estoque inicial"
        estoque.append([
            "Categoria",
            "Produto",
            "Unidade de compra",
            "Estoque mínimo",
            "Estoque inicial",
            "Observação",
        ])
        estoque.append(["Finalização", "Mousse", "frasco", 3, -1, ""])
        consumo = workbook.create_sheet("Consumo por procedimento")
        consumo.append([
            "Procedimento",
            "Categoria",
            "Produto",
            "Unidade de consumo",
            "Consumo mínimo",
            "Consumo médio",
            "Consumo máximo",
            "Observações",
        ])
        file = BytesIO()
        workbook.save(file)
        file.seek(0)

        with self.assertRaises(StockWorkbookError):
            parse_stock_workbook(file)
